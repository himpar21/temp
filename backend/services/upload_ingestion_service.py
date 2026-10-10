"""Automatic preprocessing and PostgreSQL persistence for user uploads."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import re
import tempfile
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from database.postgres.models import (
    IngestionJob,
    StoredChunk,
    StoredDocument,
    StructuredRecord,
)
from ingestion.chunking.chunk_generator import generate_chunks
from ingestion.models import SourceDocument
from ingestion.preprocessing import (
    clean_text,
    extract_entity_ids,
    infer_primary_entity_id,
    normalize_scalar,
    parse_front_matter,
    sha256_bytes,
)
from services.graph_projection_service import GraphProjectionService


SUPPORTED_EXTENSIONS = {".csv", ".json", ".jsonl", ".md", ".txt", ".pdf", ".xlsx", ".docx", ".png", ".jpg", ".jpeg"}
MAX_UPLOAD_BYTES = 25 * 1024 * 1024
logger = logging.getLogger(__name__)


def _json_hash(value: dict[str, Any]) -> str:
    encoded = json.dumps(value, ensure_ascii=False, sort_keys=True).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


class UploadIngestionService:
    def __init__(self, session: Session, *, max_chunk_words: int = 350, overlap_words: int = 50) -> None:
        self.session = session
        self.max_chunk_words = max_chunk_words
        self.overlap_words = overlap_words

    def ingest(self, filename: str, content_type: str | None, content: bytes) -> dict[str, Any]:
        safe_name = Path(filename or "upload").name
        suffix = Path(safe_name).suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file type '{suffix}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")
        if not content:
            raise ValueError("Uploaded file is empty")
        if len(content) > MAX_UPLOAD_BYTES:
            raise ValueError(f"Uploaded file exceeds the {MAX_UPLOAD_BYTES // (1024 * 1024)} MB limit")

        digest = sha256_bytes(content)
        job = IngestionJob(
            id=str(uuid.uuid4()), filename=safe_name, content_type=content_type,
            status="PROCESSING", content_hash=digest,
        )
        self.session.add(job)
        self.session.commit()

        try:
            existing = self.session.scalar(select(StoredDocument).where(StoredDocument.content_hash == digest))
            if existing:
                if suffix == ".csv":
                    reader = csv.DictReader(io.StringIO(clean_text(content.decode("utf-8-sig"))))
                    metadata = dict(existing.document_metadata or {})
                    metadata["columns"] = [
                        normalize_scalar(column) for column in (reader.fieldnames or []) if column
                    ]
                    existing.document_metadata = metadata
                graph = GraphProjectionService(self.session).project_and_persist(existing)
                try:
                    self._ensure_qdrant_index(existing)
                    job.status = "DUPLICATE"
                    job.error_message = None
                except Exception as qdrant_exc:
                    job.status = "PARTIAL"
                    job.error_message = self._qdrant_error(qdrant_exc)
                    logger.warning("Qdrant repair failed for duplicate %s: %s", safe_name, qdrant_exc)
                job.completed_at = datetime.now(timezone.utc)
                self.session.commit()
                GraphProjectionService.sync_neo4j(graph)
                return self._result(job, existing.id, duplicate=True, graph=graph)

            text, metadata, records = self._extract(safe_name, suffix, content)
            entity_ids = extract_entity_ids(text, metadata)
            logical_id = str(metadata.get("document_id") or infer_primary_entity_id(Path(safe_name), entity_ids))
            if logical_id not in entity_ids and re.fullmatch(r"(?:TP-TS|[A-Z]+)-\d+", logical_id):
                entity_ids = sorted({*entity_ids, logical_id})

            document_id = str(uuid.uuid4())
            source = SourceDocument(
                document_id=logical_id,
                source_path=f"uploads/{safe_name}",
                source_type=suffix.lstrip("."),
                document_type=str(metadata.get("document_type") or suffix.lstrip(".")),
                text=text,
                metadata={**metadata, "filename": safe_name, "content_type": content_type},
                entity_ids=entity_ids,
                content_hash=digest,
            )
            stored = StoredDocument(
                id=document_id, ingestion_job_id=job.id, logical_document_id=logical_id,
                filename=safe_name, content_type=content_type, source_type=source.source_type,
                document_type=source.document_type, text=source.text,
                document_metadata=source.metadata, entity_ids=entity_ids, content_hash=digest,
            )
            self.session.add(stored)
            self.session.flush()

            chunks = generate_chunks(source, self.max_chunk_words, self.overlap_words)
            for chunk in chunks:
                chunk_id = hashlib.sha256(f"{document_id}:{chunk.chunk_index}".encode()).hexdigest()
                self.session.add(StoredChunk(
                    id=chunk_id, document_id=document_id, chunk_index=chunk.chunk_index,
                    section=chunk.section, text=chunk.text, chunk_metadata=chunk.metadata,
                    entity_ids=chunk.entity_ids, content_hash=chunk.content_hash,
                ))

            for index, record in enumerate(records):
                self._upsert_structured_record(job.id, document_id, Path(safe_name).stem, index, record)

            graph = GraphProjectionService(self.session).project_and_persist(stored)

            # Flush first so vector failures can be recorded as retryable
            # partial ingestions without losing the PostgreSQL source data.
            self.session.flush()
            try:
                self._ensure_qdrant_index(stored, force=True)
            except Exception as qdrant_exc:
                job.status = "PARTIAL"
                job.error_message = self._qdrant_error(qdrant_exc)
                logger.warning(
                    "Qdrant upsert failed for %s (PostgreSQL data retained): %s",
                    safe_name,
                    qdrant_exc,
                )
            else:
                job.status = "COMPLETED"
                job.error_message = None
            job.document_count = 1
            job.chunk_count = len(chunks)
            job.structured_record_count = len(records)
            job.completed_at = datetime.now(timezone.utc)
            self.session.commit()
            GraphProjectionService.sync_neo4j(graph)
            return self._result(job, document_id, graph=graph)
        except Exception as exc:
            self.session.rollback()
            persisted_job = self.session.get(IngestionJob, job.id)
            if persisted_job:
                persisted_job.status = "FAILED"
                persisted_job.error_message = f"{type(exc).__name__}: {exc}"
                persisted_job.completed_at = datetime.now(timezone.utc)
                self.session.commit()
            raise

    def retry_indexing(self, job: IngestionJob) -> IngestionJob:
        """Retry the Qdrant portion of a persisted ingestion job."""
        document = self.session.scalar(
            select(StoredDocument).where(StoredDocument.content_hash == job.content_hash)
        )
        if document is None:
            raise ValueError("No persisted document is available for this ingestion job")

        try:
            chunk_count = self._ensure_qdrant_index(document, force=True)
        except Exception as exc:
            job.status = "PARTIAL"
            job.error_message = self._qdrant_error(exc)
            job.completed_at = datetime.now(timezone.utc)
            self.session.commit()
            raise RuntimeError(job.error_message) from exc

        is_original_job = document.ingestion_job_id == job.id
        job.status = "COMPLETED" if is_original_job else "DUPLICATE"
        job.error_message = None
        job.chunk_count = chunk_count
        if is_original_job:
            job.document_count = 1
        job.completed_at = datetime.now(timezone.utc)
        self.session.commit()
        return job

    def _ensure_qdrant_index(self, document: StoredDocument, *, force: bool = False) -> int:
        """Idempotently ensure every persisted chunk exists in Qdrant."""
        from database.vector.qdrant_client import get_qdrant_client

        chunks = self.session.scalars(
            select(StoredChunk)
            .where(StoredChunk.document_id == document.id)
            .order_by(StoredChunk.chunk_index)
        ).all()
        qdrant = get_qdrant_client()
        if not force and qdrant.count_document(document.id) == len(chunks):
            return len(chunks)

        payloads = [
            {
                "text": chunk.text,
                "document_id": document.id,
                "filename": document.filename,
                "entity_ids": chunk.entity_ids or [],
                "metadata": {
                    **(document.document_metadata or {}),
                    **(chunk.chunk_metadata or {}),
                    "document_type": document.document_type,
                    "section": chunk.section,
                    "chunk_index": chunk.chunk_index,
                },
            }
            for chunk in chunks
        ]
        return qdrant.upsert_chunks(payloads)

    @staticmethod
    def _qdrant_error(exc: Exception) -> str:
        return f"Qdrant indexing failed: {type(exc).__name__}: {exc}"

    def _extract(self, filename: str, suffix: str, content: bytes) -> tuple[str, dict[str, Any], list[dict[str, str]]]:
        metadata: dict[str, Any] = {"original_filename": filename}
        records: list[dict[str, str]] = []
        if suffix == ".csv":
            decoded = clean_text(content.decode("utf-8-sig"))
            reader = csv.DictReader(io.StringIO(decoded))
            columns = [normalize_scalar(column) for column in (reader.fieldnames or []) if column]
            records = self._normalize_records(reader)
            metadata.update({"document_type": "structured_csv", "columns": columns})
            return self._records_as_text(records), metadata, records
        if suffix in {".json", ".jsonl"}:
            decoded = clean_text(content.decode("utf-8-sig"))
            raw_records = [json.loads(line) for line in decoded.splitlines() if line.strip()] if suffix == ".jsonl" else json.loads(decoded)
            if isinstance(raw_records, dict):
                raw_records = [raw_records]
            if not isinstance(raw_records, list) or not all(isinstance(item, dict) for item in raw_records):
                raise ValueError("JSON upload must contain an object or an array of objects")
            records = self._normalize_records(raw_records)
            metadata.update({
                "document_type": "structured_json",
                "columns": list(records[0].keys()) if records else [],
            })
            return self._records_as_text(records), metadata, records
        if suffix == ".xlsx":
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise RuntimeError("Excel ingestion requires openpyxl") from exc
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            headers = [normalize_scalar(str(value or "")) for value in next(rows)]
            records = self._normalize_records(dict(zip(headers, row)) for row in rows)
            metadata.update({"document_type": "structured_excel", "sheet": sheet.title, "columns": headers})
            return self._records_as_text(records), metadata, records
        if suffix in {".md", ".txt"}:
            decoded = clean_text(content.decode("utf-8-sig"))
            front_matter, text = parse_front_matter(decoded) if suffix == ".md" else ({}, decoded)
            metadata.update(front_matter)
            metadata.setdefault("document_type", "markdown" if suffix == ".md" else "text")
            return text, metadata, records
        if suffix == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError as exc:
                raise RuntimeError("PDF ingestion requires pypdf") from exc
            logging.getLogger("pypdf").setLevel(logging.ERROR)
            reader = PdfReader(io.BytesIO(content), strict=False)
            text = clean_text("\n\n".join(page.extract_text() or "" for page in reader.pages))
            if not text:
                raise ValueError("PDF has no extractable text; upload OCR dependencies or a text-searchable PDF")
            metadata.update({"document_type": "pdf", "page_count": len(reader.pages), "ocr_required": False})
            return text, metadata, records
        if suffix == ".docx":
            try:
                from docx import Document
            except ImportError as exc:
                raise RuntimeError("DOCX ingestion requires python-docx") from exc
            document = Document(io.BytesIO(content))
            metadata["document_type"] = "docx"
            return clean_text("\n".join(paragraph.text for paragraph in document.paragraphs)), metadata, records
        if suffix in {".png", ".jpg", ".jpeg"}:
            try:
                import pytesseract
                from PIL import Image
            except ImportError as exc:
                raise RuntimeError("Image OCR requires Pillow and pytesseract") from exc
            text = clean_text(pytesseract.image_to_string(Image.open(io.BytesIO(content))))
            if not text:
                raise ValueError("OCR did not extract text from the uploaded image")
            metadata.update({"document_type": "image", "ocr_required": True, "ocr_applied": True})
            return text, metadata, records
        raise ValueError(f"Unsupported file type: {suffix}")

    @staticmethod
    def _normalize_records(rows: Any) -> list[dict[str, str]]:
        normalized: list[dict[str, str]] = []
        seen: set[str] = set()
        for row in rows:
            record = {normalize_scalar(str(key)): normalize_scalar(str(value or "")) for key, value in dict(row).items() if key is not None}
            if not record:
                continue
            digest = _json_hash(record)
            if digest not in seen:
                seen.add(digest)
                normalized.append(record)
        return normalized

    @staticmethod
    def _records_as_text(records: list[dict[str, str]]) -> str:
        return "\n".join(json.dumps(record, ensure_ascii=False, sort_keys=True) for record in records)

    def _upsert_structured_record(self, job_id: str, document_id: str, record_type: str, index: int, record: dict[str, str]) -> None:
        values = list(record.values())
        record_key = values[0] if values and values[0] else f"row-{index + 1}"
        existing = self.session.scalar(select(StructuredRecord).where(
            StructuredRecord.record_type == record_type,
            StructuredRecord.record_key == record_key,
        ))
        if existing:
            existing.ingestion_job_id = job_id
            existing.document_id = document_id
            existing.payload = record
            existing.content_hash = _json_hash(record)
        else:
            self.session.add(StructuredRecord(
                ingestion_job_id=job_id, document_id=document_id, record_type=record_type,
                record_key=record_key, payload=record, content_hash=_json_hash(record),
            ))

    @staticmethod
    def _result(
        job: IngestionJob,
        document_id: str,
        duplicate: bool = False,
        graph: Any | None = None,
    ) -> dict[str, Any]:
        return {
            "job_id": job.id, "status": job.status, "filename": job.filename,
            "duplicate": duplicate, "document_id": document_id,
            "structured_records": job.structured_record_count, "chunks": job.chunk_count,
            "graph_nodes": len(graph.nodes) if graph else 0,
            "graph_edges": len(graph.edges) if graph else 0,
            "error_message": job.error_message,
        }
